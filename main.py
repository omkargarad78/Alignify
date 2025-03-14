from flask import Flask, request, render_template, jsonify
from werkzeug.utils import secure_filename
import os
import PyPDF2
import docx
from sklearn.feature_extraction.text import TfidfVectorizer, CountVectorizer
from sklearn.metrics.pairwise import cosine_similarity
import re
import nltk
from nltk.corpus import stopwords
from collections import Counter
import json
import google.generativeai as genai
from nltk.tokenize import word_tokenize
import string
import numpy as np
from dotenv import load_dotenv

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads'

# Custom JSON encoder to handle any non-serializable data
class CustomJSONEncoder(json.JSONEncoder):
    def default(self, obj):
        try:
            return super().default(obj)
        except TypeError:
            return str(obj)

app.json_encoder = CustomJSONEncoder

load_dotenv()
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
genai.configure(api_key=GEMINI_API_KEY)

# Available Gemini models (use models that actually exist)
GEMINI_MODEL = "gemini-1.5-flash"  # Updated model name

# Ensure the upload folder exists
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

# Download NLTK resources if not available
try:
    nltk.download('punkt')
    nltk.download('stopwords')
except Exception as e:
    print(f"Error downloading NLTK resources: {e}")

# List of technical terms to focus on when extracting keywords
TECH_TERMS = set([
    "python", "java", "javascript", "typescript", "c++", "c#", "go", "rust", "php", "ruby", "swift",
    "kotlin", "scala", "perl", "bash", "shell", "sql", "nosql", "react", "angular", "vue", "node", 
    "express", "django", "flask", "fastapi", "spring", "laravel", "rails", "asp.net", "dotnet",
    "aws", "azure", "gcp", "docker", "kubernetes", "terraform", "ansible", "jenkins", "circleci", 
    "git", "github", "gitlab", "bitbucket", "jira", "confluence", "mongodb", "mysql", "postgresql", 
    "oracle", "sqlserver", "redis", "cassandra", "elasticsearch", "dynamodb", "kafka", "rabbitmq", 
    "tensorflow", "pytorch", "scikit-learn", "pandas", "numpy", "matplotlib", "seaborn", "hadoop", 
    "spark", "airflow", "linux", "windows", "macos", "ios", "android", "react native", "flutter", 
    "xamarin", "unity", "unreal", "graphql", "rest", "soap", "microservices", "serverless", "agile", 
    "scrum", "kanban", "devops", "mlops", "machine learning", "deep learning", "ai", "blockchain", 
    "iot", "websocket", "websockets", "webrtc", "css", "html", "sass", "less", "tailwind", "bootstrap", 
    "material ui", "figma", "sketch", "adobe xd", "photoshop", "illustrator", "npm", "yarn", "webpack", 
    "babel", "vite", "jest", "mocha", "cypress", "selenium", "pytest", "junit", "oauth", "jwt", "api", "sdk",
    "saas", "paas", "iaas", "ci/cd", "etl", "data science", "data engineering", "web development", 
    "mobile development", "full stack", "frontend", "backend", "qa", "testing", "security", "uiux", 
    "ui", "ux", "product management", "project management", "hci", "computer vision", "nlp", "networking",
    "zmq", "zeromq", "0mq", "postgres", "postgresql", "real-time", "realtime", "trading", "visualization",
    "high-performance", "nextjs", "next.js", "financial", "containerization", "containers", "web socket",
    "web-socket", "cloud computing", "cloud", "database"
])

# Function to extract text from PDF
def extract_text_from_pdf(pdf_path):
    text = ''
    try:
        with open(pdf_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            for page in pdf_reader.pages:
                text += page.extract_text() if page.extract_text() else ''
        return text
    except Exception as e:
        print(f"Error extracting text from PDF: {e}")
        return text

# Function to extract text from DOCX
def extract_text_from_docx(docx_path):
    try:
        doc = docx.Document(docx_path)
        return '\n'.join([para.text for para in doc.paragraphs])
    except Exception as e:
        print(f"Error extracting text from DOCX: {e}")
        return ""

# Function to calculate resume matching scores
def calculate_matching_scores(job_description, resumes):
    """
    Improved matching algorithm that focuses on technical skills and important keywords
    rather than general text similarity.
    """
    scores = []
    
    for resume_text in resumes:
        # Extract technical keywords from both documents
        jd_keywords = extract_all_jd_keywords(job_description)
        resume_keywords = extract_all_jd_keywords(resume_text)  # Using the same extraction for resume
        
        # Calculate matches based on important technical skills
        # 1. Calculate keyword overlap score
        if jd_keywords:
            matched_keywords = [k for k in jd_keywords if k in resume_text.lower()]
            keyword_match_score = len(matched_keywords) / len(jd_keywords) if jd_keywords else 0
        else:
            keyword_match_score = 0
            
        # 2. Use TF-IDF for semantic similarity but weight it less
        documents = [job_description, resume_text]
        vectorizer = TfidfVectorizer(stop_words='english')
        try:
            # Try to create TF-IDF vectors
            vectors = vectorizer.fit_transform(documents)
            tfidf_score = cosine_similarity(vectors[0:1], vectors[1:2]).flatten()[0]
        except:
            # Fallback if vectorization fails
            tfidf_score = 0.5  # Neutral score
        
        # 3. Calculate specific role and skill matches from qualifications section
        # Extract qualifications section from JD if it exists
        qual_match_score = 0
        if "qualifications:" in job_description.lower():
            qualifications = job_description.lower().split("qualifications:")[1].split("\n\n")[0]
            qual_lines = [q.strip('- ').lower() for q in qualifications.split('\n') if q.strip()]
            
            # Count how many qualification lines have matching content in resume
            matched_quals = 0
            for qual in qual_lines:
                # Check if key parts of the qualification are in the resume
                if any(key_part in resume_text.lower() for key_part in qual.split()):
                    matched_quals += 1
            
            qual_match_score = matched_quals / len(qual_lines) if qual_lines else 0
        
        # Weight the different components of the score
        # Technical keywords are most important (60%)
        # Qualifications match is important (25%)
        # TF-IDF similarity provides context (15%)
        final_score = (keyword_match_score * 0.6) + (qual_match_score * 0.25) + (tfidf_score * 0.15)
        
        # Normalize to ensure scores reflect expectations better (boost scores above 50%)
        # This better represents when a resume matches more than half the requirements
        if final_score > 0.5:
            normalized_score = 0.5 + (final_score - 0.5) * 1.5
        else:
            normalized_score = final_score
        
        # Ensure score is within 0-1 range
        normalized_score = max(0, min(1, normalized_score))
        
        scores.append(normalized_score)
    
    return np.array(scores)

# Function to extract keywords from text, focusing on technical terms
def extract_keywords(text, top_n=30, focus_on_tech=True):
    try:
        # Clean the text and convert to lowercase
        text = text.lower()
        # Remove punctuation
        text = text.translate(str.maketrans('', '', string.punctuation))
        # Split into words
        words = text.split()
        
        # Filter out stopwords
        stop_words = set(stopwords.words('english'))
        words = [word for word in words if word not in stop_words and len(word) > 2]
        
        # Focus on technical terms if required
        if focus_on_tech:
            # Check for tech terms and common variations
            tech_words = []
            for word in words:
                if word in TECH_TERMS:
                    tech_words.append(word)
                # Check for compound terms like "web socket" by combining with next word
                for i in range(len(words) - 1):
                    compound = words[i] + " " + words[i+1]
                    if compound in TECH_TERMS:
                        tech_words.append(compound)
            words = tech_words
        
        # Count word frequencies
        word_freq = Counter(words)
        
        # Return most common words
        return [word for word, _ in word_freq.most_common(top_n)]
    except Exception as e:
        print(f"Error extracting keywords: {e}")
        return []

# Function to extract all keywords from a text, not just job description
def extract_all_keywords(text):
    """Extract all technical and important keywords from any text including resumes"""
    # Convert to lowercase and clean
    text = text.lower()
    
    keywords = []
    
    # Check for all tech terms in the text
    for term in TECH_TERMS:
        # Look for exact matches or variations
        if term in text or term.replace('-', ' ') in text or term.replace(' ', '-') in text:
            keywords.append(term)
    
    # Also check for other important professional terms
    professional_terms = [
        'degree', 'bachelor', 'master', 'phd', 'certified', 'certification',
        'experience', 'year', 'senior', 'junior', 'lead', 'manager', 'engineer',
        'developer', 'professional', 'analyst', 'specialist', 'consultant',
        'project', 'team', 'collaborate', 'leadership', 'initiative'
    ]
    
    for term in professional_terms:
        if term in text:
            keywords.append(term)
    
    # Also check for specific mentioned technologies using regex
    tech_patterns = [
        r'experience (?:with|in) ([\w\s\-\+]+)',
        r'knowledge of ([\w\s\-\+]+)',
        r'familiar(?:ity)? with ([\w\s\-\+]+)',
        r'proficien(?:t|cy) (?:with|in) ([\w\s\-\+]+)',
        r'skill(?:s|ed)(?::|in) ([\w\s\-\+,]+)',
        r'technolog(?:y|ies)(?::|in) ([\w\s\-\+,]+)',
        r'using ([\w\s\-\+]+)',
        r'work(?:ed|ing) with ([\w\s\-\+]+)',
        r'built (?:with|using) ([\w\s\-\+]+)',
        r'developed (?:with|using|in) ([\w\s\-\+]+)'
    ]
    
    for pattern in tech_patterns:
        matches = re.findall(pattern, text)
        for match in matches:
            # Split by commas or "and" to get individual technologies
            technologies = re.split(r',|\sand\s', match)
            for tech in technologies:
                tech = tech.strip().lower()
                if tech and len(tech) > 2:  # Avoid very short terms
                    keywords.append(tech)
    
    # Return unique keywords
    return list(set(keywords))

# Update the existing extract_all_jd_keywords function to use the new more comprehensive function
def extract_all_jd_keywords(job_description):
    return extract_all_keywords(job_description)

# Function to find missing keywords in resume, focusing on technical terms
def find_missing_keywords(job_description, resume_text):
    # Get all technical keywords from job description
    jd_keywords = extract_all_jd_keywords(job_description)
    
    # Convert resume text to lowercase for case-insensitive matching
    resume_text = resume_text.lower()
    
    # Create a list of missing keywords
    missing = []
    
    # Process technical keywords specifically mentioned in job description
    for keyword in jd_keywords:
        # Handle variations for technologies
        if keyword in ["websocket", "websockets", "web socket", "web-socket"]:
            if not any(term in resume_text for term in ["websocket", "websockets", "web socket", "web-socket"]):
                missing.append("WebSocket")
        elif keyword in ["zmq", "zeromq", "0mq"]:
            if not any(term in resume_text for term in ["zmq", "zeromq", "0mq"]):
                missing.append("ZMQ")
        elif keyword in ["c++", "c plus plus"]:
            if not any(term in resume_text for term in ["c++", "c plus plus"]):
                missing.append("C++")
        elif keyword in ["postgres", "postgresql"]:
            if not any(term in resume_text for term in ["postgres", "postgresql"]):
                missing.append("PostgreSQL")
        elif keyword in ["real-time", "realtime", "real time"]:
            if not any(term in resume_text for term in ["real-time", "realtime", "real time"]):
                missing.append("Real-time processing")
        elif keyword in ["react", "reactjs"]:
            if not any(term in resume_text for term in ["react", "reactjs"]):
                missing.append("React")
        # For other technologies, do a simple check
        elif keyword not in resume_text:
            # Capitalize first letter of each word for display
            display_keyword = ' '.join(word.capitalize() for word in keyword.split())
            missing.append(display_keyword)
    
    # Return unique keywords, properly formatted
    return list(set(missing))

# Function to generate personalized email template using Gemini
def generate_email_template_with_gemini(job_description, resume_text, match_score):
    try:
        # Extract company name if possible
        company_match = re.search(r'(company|organization|firm)[:\s]+([\w\s]+)', job_description, re.IGNORECASE)
        company = company_match.group(2).strip() if company_match else "the company"
        
        # Extract company description if possible
        company_description_match = re.search(r'((?:we are|our company|about us)[^.!?]*)', job_description, re.IGNORECASE)
        company_description = company_description_match.group(1).strip() if company_description_match else None
        
        # Extract role if possible
        role_match = re.search(r'(position|role|job title|engineer)[:\s]+([\w\s]+)', job_description, re.IGNORECASE)
        role = role_match.group(2).strip() if role_match else "Software Engineer"
        
        # Extract skills and requirements from both documents
        jd_keywords = extract_all_jd_keywords(job_description)
        resume_keywords = extract_all_keywords(resume_text)
        shared_skills = [word for word in resume_keywords if word.lower() in [k.lower() for k in jd_keywords]]
        
        # Extract education and experience from resume
        education_match = re.search(r'(?:education|university|college|degree)([^.]*.(?:degree|bachelor|master|phd)[^.]*\.)', resume_text, re.IGNORECASE)
        education = education_match.group(1).strip() if education_match else None
        
        # Create prompt for Gemini
        prompt = f"""
        Create a personalized job application email for the {role} position at {company}. The email should be tailored based on the following information:

        JOB DESCRIPTION SUMMARY:
        {job_description[:500]}...

        RESUME SUMMARY:
        {resume_text[:500]}...

        MATCHING SKILLS BETWEEN RESUME AND JOB:
        {', '.join(shared_skills[:10]) if shared_skills else "No direct skill matches found"}

        IMPORTANT REQUIREMENTS:
        1. Create a completely original email structure based on this specific job and resume - DO NOT follow a template (Strictly small up to 1-2 short paragraphs)
        2. Only mention skills and experience that appear in the resume,as bullet points
        3. Highlight specific connections between the candidate's experience and the company's needs
        4. Maintain a professional but conversational tone
        5. Include a brief introduction, 2-3 relevant qualifications/experiences, and a closing paragraph
        6. Ensure the email feels authentically tailored to THIS specific job, not generic
        
        7. Don't include generic "fill in the blank" placeholders that the user must complete
        8. Vary sentence structure and paragraph flow to feel natural and conversational
        9. Create a dynamic subject line that references both the position and a key qualification
        10. At the end, below best regards, add the person's name and if the linkedin or github profile are there or other also inculde them one blow the other 
        11. I don't want to fill in the blanks, so don't include any placeholders
        """
        
        # Generate content with Gemini
        try:
            model = genai.GenerativeModel(GEMINI_MODEL)
            response = model.generate_content(prompt)
            return response.text
        except Exception as gemini_err:
            print(f"Specific Gemini error: {gemini_err}")
            raise  # Re-raise to be caught by the outer exception handler
        
    except Exception as e:
        print(f"Error generating email with Gemini: {e}")
        
        # Fallback to a simple template if Gemini fails
        shared_skills_str = ', '.join(extract_keywords(resume_text, 5, True)[:2])
        
        return f"""Subject: Application for {role} Position

Dear Hiring Team at {company},

I am writing to express my interest in the {role} position. Based on the job description, I believe my background in {shared_skills_str} would be valuable to your team.

[Personalize this section with your relevant experience and qualifications]

I look forward to the opportunity to discuss how my skills align with your needs.

Best regards,
[Your Name]
"""

# Function to generate LinkedIn message using Gemini
def generate_linkedin_message_with_gemini(job_description, resume_text, match_score):
    try:
        # Extract key information
        jd_keywords = extract_all_jd_keywords(job_description)
        resume_keywords = extract_all_keywords(resume_text)
        shared_skills = [word for word in resume_keywords if word.lower() in [k.lower() for k in jd_keywords]]
        
        # Extract role if possible
        role_match = re.search(r'(position|role|job title|engineer)[:\s]+([\w\s]+)', job_description, re.IGNORECASE)
        role = role_match.group(2).strip() if role_match else "Software Engineer"
        
        # Create prompt for Gemini
        prompt = f"""
        Create a very brief LinkedIn message to any employee at a company about a job application. 
        
        JOB ROLE: {role}
        TOP MATCHING SKILL: {shared_skills[0] if shared_skills else "technical"}
        
        REQUIREMENTS:
        1. Start with "Hi [Name]," - this is the ONLY placeholder allowed
        2. STRICT 200 character limit MAXIMUM including the placeholder
        3. Make it clear you're applying for the specific role
        4. Also tell what the role i am looking for
        5. Include a brief question about the application process or role
        6. Create a message that would make sense to ANY employee (not just HR)
        7. Make it immediately clear what you're requesting
        8. Extremely concise but still conversational and professional
        9. If possible also ask for refferal, or any guide from the person
        """
        
        # Generate content with Gemini
        try:
            model = genai.GenerativeModel(GEMINI_MODEL)
            response = model.generate_content(prompt)
            message = response.text.strip()
            
            # Ensure the message is under 200 characters
            if len(message) > 200:
                # Try to make a clean cutoff
                return message[:197] + "..."
            return message
        except Exception as gemini_err:
            print(f"Specific Gemini error: {gemini_err}")
            raise  # Re-raise to be caught by the outer exception handler
        
    except Exception as e:
        print(f"Error generating LinkedIn message with Gemini: {e}")
        
        # Fallback to a simple template
        skill = extract_keywords(resume_text, 1, True)[0] if extract_keywords(resume_text, 1, True) else "technical"
        
        return f"""Hi [Name], I've applied for the {role} position at your company. My {skill} experience seems like a great fit. Could you share any insights about the team or hiring process?"""

# Function to generate resume improvement suggestions using Gemini
def generate_improvement_suggestions_with_gemini(job_description, resume_text):
    try:
        # Extract technical keywords and find missing ones
        missing_keywords = find_missing_keywords(job_description, resume_text)
        
        # Create prompt for Gemini
        prompt = f"""
        Generate specific, actionable resume improvement suggestions for a technical role with the following details:
        
        Job Description Summary: {job_description[:300]}...
        
        Missing Technical Skills/Keywords: {', '.join(missing_keywords[:10])}
        
        Requirements:
        1. Provide 4-5 specific, actionable suggestions focused on technical aspects
        2. Each suggestion should be 1-2 sentences and very specific
        3. Focus mainly on how to address the missing technical skills
        4. Include advice about formatting and quantifying achievements
        5. Be direct and practical, not generic
        6. Focus only on technical/professional aspects, not soft skills
        7. Ensure perfect grammar throughout
        """
        
        # Generate content with Gemini
        try:
            model = genai.GenerativeModel(GEMINI_MODEL)
            response = model.generate_content(prompt)
            
            # Process the response to get a list of suggestions
            suggestions_text = response.text.strip()
            suggestions_list = [s.strip() for s in suggestions_text.split('\n') if s.strip()]
            
            # Extract just the first 5 suggestions, remove any numbering or bullet points
            cleaned_suggestions = []
            for suggestion in suggestions_list[:5]:
                # Remove numbers, dashes, asterisks at the beginning
                cleaned = re.sub(r'^[\d\.\-\*\•]+\s*', '', suggestion)
                cleaned_suggestions.append(cleaned)
            
            return cleaned_suggestions, missing_keywords
        except Exception as gemini_err:
            print(f"Specific Gemini error: {gemini_err}")
            raise  # Re-raise to be caught by the outer exception handler
        
    except Exception as e:
        print(f"Error generating suggestions with Gemini: {e}")
        
        # Fallback to more specific technical suggestions
        tech_suggestions = [
            f"Add experience with {', '.join(missing_keywords[:3])} to your technical skills section.",
            f"Include a project that demonstrates your ability to work with real-time data processing or trading systems.",
            "Quantify your achievements with specific metrics (e.g., 'Reduced API response time by 30%').",
            "Add links to your GitHub repositories that showcase relevant code samples.",
            "List any relevant technical courses or certifications related to financial technology."
        ]
        
        return tech_suggestions, missing_keywords

# Route for the frontend
@app.route('/')
def index():
    return render_template('index.html')

# Route to handle file uploads and matching
@app.route('/match', methods=['POST'])
def match_resumes():
    try:
        job_description = request.form['job_description']
        uploaded_files = request.files.getlist('resumes')

        resumes = []
        filenames = []
        for file in uploaded_files:
            if file and (file.filename.endswith('.pdf') or file.filename.endswith('.docx')):
                filename = secure_filename(file.filename)
                filenames.append(filename)
                file_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
                file.save(file_path)
                if filename.endswith('.pdf'):
                    resumes.append(extract_text_from_pdf(file_path))
                elif filename.endswith('.docx'):
                    resumes.append(extract_text_from_docx(file_path))

        if not resumes:
            return jsonify({'error': 'No valid resume files uploaded'}), 400

        scores = calculate_matching_scores(job_description, resumes)

        return jsonify({'scores': scores.tolist(), 'filenames': filenames})
    except Exception as e:
        print(f"Error in match_resumes: {e}")
        return jsonify({'error': str(e)}), 500

# New route for resume analysis
@app.route('/analyze', methods=['POST'])
def analyze_resume():
    try:
        # Get job description directly from the form
        job_description = request.form.get('job_description', '')
        
        # Get resume file
        resume_file = request.files.get('resume')
        
        if not job_description or not resume_file:
            return jsonify({'error': 'Missing job description or resume'}), 400
        
        # Save and process the resume file
        filename = secure_filename(resume_file.filename)
        file_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        resume_file.save(file_path)
        
        # Extract text from the resume
        if filename.endswith('.pdf'):
            resume_text = extract_text_from_pdf(file_path)
        elif filename.endswith('.docx'):
            resume_text = extract_text_from_docx(file_path)
        else:
            return jsonify({'error': 'Unsupported file format'}), 400
        
        if not resume_text:
            return jsonify({'error': 'Could not extract text from resume'}), 400
        
        # Calculate match score using improved algorithm
        match_score = calculate_matching_scores(job_description, [resume_text])[0]
        
        # Extract keywords from both for display and debugging
        jd_keywords = extract_all_jd_keywords(job_description)
        resume_keywords = extract_all_keywords(resume_text)
        matching_keywords = [k for k in jd_keywords if k in resume_text.lower()]
        
        # Generate email and LinkedIn templates using Gemini
        email_template = generate_email_template_with_gemini(job_description, resume_text, match_score)
        linkedin_message = generate_linkedin_message_with_gemini(job_description, resume_text, match_score)
        
        # Generate improvement suggestions using Gemini
        suggestions, missing_keywords = generate_improvement_suggestions_with_gemini(job_description, resume_text)
        
        # Prepare and return response with extra diagnostic info
        return jsonify({
            'match_score': match_score,
            'email_template': email_template,
            'linkedin_message': linkedin_message,
            'improvement_suggestions': suggestions,
            'missing_keywords': missing_keywords,
            'debug_info': {
                'jd_keywords_count': len(jd_keywords),
                'resume_keywords_count': len(resume_keywords),
                'matching_keywords_count': len(matching_keywords),
                'matching_keywords': matching_keywords[:15]  # Show first 15 matches for debugging
            }
        })
    except Exception as e:
        print(f"Error in analyze_resume: {e}")
        return jsonify({'error': str(e)}), 500

if __name__ == '__main__':
    app.run(debug=True, port=8000)
